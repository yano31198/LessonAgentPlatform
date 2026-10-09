from __future__ import annotations

import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Iterable

from docx import Document


def _round(value: float) -> float:
    return round(float(value), 2)


def build_comparison_payload(results: list[dict]) -> dict:
    """Build a model-free comparison payload from completed F1 run results."""
    if len(results) < 2:
        raise ValueError("至少需要 2 个已完成的 F1 结果进行对比")

    items = []
    dimension_names: list[str] = []
    group_names: list[str] = []

    for result in results:
        rubric = result.get("rubric") or {}
        dims = rubric.get("dimensions") or []
        groups = rubric.get("groups") or []
        for dim in dims:
            name = dim.get("name")
            if name and name not in dimension_names:
                dimension_names.append(name)
        for group in groups:
            name = group.get("name")
            if name and name not in group_names:
                group_names.append(name)

        items.append(
            {
                "runId": result.get("runId"),
                "lessonId": result.get("lessonId"),
                "versionId": result.get("versionId"),
                "subject": result.get("subject"),
                "grade": result.get("grade"),
                "topic": result.get("topic"),
                "score": (result.get("score") or {}).get("total", rubric.get("score")),
                "maximum": (result.get("score") or {}).get("maximum", rubric.get("maximum", 100)),
            }
        )

    dimensions = []
    for name in dimension_names:
        rows = []
        for result in results:
            dim = next((d for d in (result.get("rubric") or {}).get("dimensions", []) if d.get("name") == name), None)
            if dim:
                rows.append(
                    {
                        "runId": result.get("runId"),
                        "lessonId": result.get("lessonId"),
                        "versionId": result.get("versionId"),
                        "score": dim.get("score"),
                        "maximum": dim.get("maximum"),
                        "grade": dim.get("grade"),
                    }
                )
        dimensions.append({"name": name, "items": rows})

    groups_payload = []
    for name in group_names:
        rows = []
        for result in results:
            group = next((g for g in (result.get("rubric") or {}).get("groups", []) if g.get("name") == name), None)
            if group:
                rows.append(
                    {
                        "runId": result.get("runId"),
                        "lessonId": result.get("lessonId"),
                        "versionId": result.get("versionId"),
                        "score": group.get("score"),
                        "maximum": group.get("maximum"),
                    }
                )
        groups_payload.append({"name": name, "items": rows})

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "count": len(results),
        "items": items,
        "groups": groups_payload,
        "dimensions": dimensions,
    }


def build_batch_quality_payload(results: list[dict]) -> dict:
    """Aggregate multiple completed F1 results into a deterministic quality report."""
    if not results:
        raise ValueError("没有可用于生成批量质量报告的 F1 结果")

    lesson_rows = []
    scores: list[float] = []
    dim_acc: dict[str, dict] = {}

    for result in results:
        rubric = result.get("rubric") or {}
        score = float((result.get("score") or {}).get("total", rubric.get("score", 0)) or 0)
        maximum = float((result.get("score") or {}).get("maximum", rubric.get("maximum", 100)) or 100)
        scores.append(score)
        lesson_rows.append(
            {
                "runId": result.get("runId"),
                "lessonId": result.get("lessonId"),
                "versionId": result.get("versionId"),
                "subject": result.get("subject"),
                "grade": result.get("grade"),
                "topic": result.get("topic") or result.get("lessonId"),
                "score": _round(score),
                "maximum": _round(maximum),
                "summary": result.get("summary", ""),
            }
        )

        raw_rows = {
            row.get("维度名称"): row
            for row in (result.get("rawResult") or {}).get("维度评分", [])
            if isinstance(row, dict) and row.get("维度名称")
        }
        for dim in rubric.get("dimensions", []) or []:
            name = dim.get("name")
            if not name:
                continue
            cap = float(dim.get("maximum") or 0)
            value = float(dim.get("score") or 0)
            grade = dim.get("grade") or ""
            entry = dim_acc.setdefault(
                name,
                {
                    "scores": [],
                    "percentages": [],
                    "grades": Counter(),
                    "annotationCount": 0,
                    "sampleProblems": [],
                    "sampleRecommendations": [],
                    "maximum": cap,
                },
            )
            entry["scores"].append(value)
            entry["percentages"].append(value / cap if cap else 0.0)
            if grade:
                entry["grades"][grade] += 1

            raw = raw_rows.get(name) or {}
            annotations = raw.get("批注列表") or []
            entry["annotationCount"] += len(annotations)
            for ann in annotations:
                if not isinstance(ann, dict):
                    continue
                problem = ann.get("具体分析") or ann.get("建议")
                recommendation = ann.get("细化建议") or ann.get("建议")
                if problem and problem not in entry["sampleProblems"] and len(entry["sampleProblems"]) < 3:
                    entry["sampleProblems"].append(problem)
                if recommendation and recommendation not in entry["sampleRecommendations"] and len(entry["sampleRecommendations"]) < 3:
                    entry["sampleRecommendations"].append(recommendation)

    lesson_rows.sort(key=lambda x: x["score"], reverse=True)
    for idx, row in enumerate(lesson_rows, start=1):
        row["rank"] = idx

    dimension_rows = []
    for name, data in dim_acc.items():
        average = mean(data["scores"]) if data["scores"] else 0.0
        avg_pct = mean(data["percentages"]) if data["percentages"] else 0.0
        dimension_rows.append(
            {
                "name": name,
                "averageScore": _round(average),
                "maximum": _round(data["maximum"]),
                "averageRate": _round(avg_pct * 100),
                "gradeDistribution": {k: int(data["grades"].get(k, 0)) for k in ("A", "B", "C")},
                "annotationCount": int(data["annotationCount"]),
                "sampleProblems": data["sampleProblems"],
                "sampleRecommendations": data["sampleRecommendations"],
            }
        )

    ranked_dims = sorted(dimension_rows, key=lambda x: (x["averageRate"], -x["annotationCount"]), reverse=True)
    strengths = [
        {
            "dimension": x["name"],
            "averageRate": x["averageRate"],
            "averageScore": x["averageScore"],
            "maximum": x["maximum"],
            "gradeDistribution": x["gradeDistribution"],
        }
        for x in ranked_dims[:3]
    ]
    weaknesses_src = sorted(dimension_rows, key=lambda x: (x["averageRate"], -x["gradeDistribution"].get("C", 0), -x["annotationCount"]))[:3]
    common_issues = [
        {
            "dimension": x["name"],
            "averageRate": x["averageRate"],
            "cCount": x["gradeDistribution"].get("C", 0),
            "annotationCount": x["annotationCount"],
            "problemExamples": x["sampleProblems"],
            "recommendationExamples": x["sampleRecommendations"],
        }
        for x in weaknesses_src
    ]

    return {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "summary": {
            "lessonCount": len(results),
            "averageScore": _round(mean(scores)),
            "highestScore": _round(max(scores)),
            "lowestScore": _round(min(scores)),
            "maximum": 100,
        },
        "strengths": strengths,
        "commonIssues": common_issues,
        "dimensions": dimension_rows,
        "lessons": lesson_rows,
        "generationMethod": "rule-based aggregation of completed F1 results; no additional LLM call",
    }


def write_batch_quality_report(results: list[dict], output_dir: str | Path) -> dict[str, str]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_batch_quality_payload(results)

    json_path = output_dir / "batch_quality_report.json"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    doc = Document()
    doc.add_heading("批量教案质量综合报告", level=0)
    summary = payload["summary"]
    doc.add_paragraph(
        f"本批次共评审 {summary['lessonCount']} 份教案；平均分 {summary['averageScore']}/{summary['maximum']}，"
        f"最高分 {summary['highestScore']}，最低分 {summary['lowestScore']}。"
    )

    doc.add_heading("一、批次概况", level=1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["教案/版本", "主题", "学科", "得分", "排名"]
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for row in payload["lessons"]:
        cells = table.add_row().cells
        cells[0].text = f"{row.get('lessonId') or ''} / {row.get('versionId') or ''}"
        cells[1].text = str(row.get("topic") or "")
        cells[2].text = str(row.get("subject") or "")
        cells[3].text = f"{row['score']}/{row['maximum']}"
        cells[4].text = str(row["rank"])

    doc.add_heading("二、整体优势", level=1)
    for item in payload["strengths"]:
        dist = item["gradeDistribution"]
        doc.add_paragraph(
            f"{item['dimension']}：平均得分率 {item['averageRate']}%，"
            f"A/B/C 分布 {dist['A']}/{dist['B']}/{dist['C']}。",
            style="List Bullet",
        )

    doc.add_heading("三、共性问题与改进方向", level=1)
    for item in payload["commonIssues"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{item['dimension']}：").bold = True
        p.add_run(
            f"平均得分率 {item['averageRate']}%，C 等级 {item['cCount']} 次，"
            f"累计批注 {item['annotationCount']} 条。"
        )
        if item["problemExamples"]:
            doc.add_paragraph("典型诊断：" + item["problemExamples"][0])
        if item["recommendationExamples"]:
            doc.add_paragraph("改进建议：" + item["recommendationExamples"][0])

    doc.add_heading("四、十二维度整体表现", level=1)
    dim_table = doc.add_table(rows=1, cols=6)
    headers = ["维度", "平均分", "满分", "得分率", "A/B/C", "批注数"]
    for i, value in enumerate(headers):
        dim_table.rows[0].cells[i].text = value
    for row in payload["dimensions"]:
        cells = dim_table.add_row().cells
        cells[0].text = row["name"]
        cells[1].text = str(row["averageScore"])
        cells[2].text = str(row["maximum"])
        cells[3].text = f"{row['averageRate']}%"
        dist = row["gradeDistribution"]
        cells[4].text = f"{dist['A']}/{dist['B']}/{dist['C']}"
        cells[5].text = str(row["annotationCount"])

    doc.add_heading("五、说明", level=1)
    doc.add_paragraph(
        "本综合报告由已完成的 F1 单份评分结果自动聚合生成，不会再次调用大模型。"
        "原始逐份评分、批注与下载产物仍以各自 runId 为准。"
    )

    docx_path = output_dir / "batch_quality_report.docx"
    doc.save(docx_path)
    return {"quality_json": str(json_path), "quality_docx": str(docx_path)}
