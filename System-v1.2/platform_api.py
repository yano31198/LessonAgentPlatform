"""F1 的本机 HTTP 适配层；原 Streamlit 入口保持独立运行。

运行：python -m uvicorn platform_api:app --host 127.0.0.1 --port 8002 --http h11
平台仅传已经保存的教案版本文本。每个任务使用独立输入/输出目录。

F1 平台固定口径：
- criteria_type = "student"
- variant = "v2_guided"
- suggest 默认开启，可通过请求字段临时关闭
"""
from __future__ import annotations

import json
import logging
import re
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fastapi import FastAPI, HTTPException, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, Field

from pipeline import load_config
from providers import get_key_status
from service import run_annotation
from standards import score_to_grade
from comparison import generate_comparison
from batch_reporting import build_comparison_payload, write_batch_quality_report

ROOT = Path(__file__).resolve().parent
RUN_ROOT = ROOT / "platform_runs"
BATCH_ROOT = ROOT / "platform_batches"
COMPARISON_ROOT = ROOT / "platform_comparisons"
EXECUTOR = ThreadPoolExecutor(max_workers=1, thread_name_prefix="f1-annotation")
LOCK = threading.Lock()
LOGGER = logging.getLogger(__name__)

CRITERIA_TYPE = "student"
PROMPT_VARIANT = "v2_guided"
DEFAULT_SUGGEST = True

# 一级维度来自 F1 的 student_评价指标V3.docx；细项名称来自 prompts.STUDENT_DIMS。
GROUPS = (
    ("思想理念", 10, ("教学理念", "课程思政")),
    ("学情与目标分析", 10, ("学情分析", "教学目标")),
    ("内容与方法", 30, ("教学内容", "教学资源", "教学方法")),
    ("过程与实施", 30, ("过程设计", "课堂互动")),
    ("评价反馈", 10, ("评价反馈",)),
    ("文档规范", 10, ("版式观感", "结构完整")),
)
EXPECTED = {name for _, _, children in GROUPS for name in children}

app = FastAPI(title="F1 Annotation Adapter", docs_url=None, redoc_url=None)


@app.exception_handler(RequestValidationError)
async def validation_error(_request: Request, error: RequestValidationError):
    """避免把完整教案正文写进校验错误日志。"""
    fields = [
        ".".join(map(str, item.get("loc", ()))) + ": " + item.get("type", "unknown")
        for item in error.errors()
    ]
    body = error.body
    received_bytes = len(body) if isinstance(body, bytes) else (
        len(body.encode("utf-8")) if isinstance(body, str) else 0
    )
    LOGGER.warning(
        "F1 request validation failed: %s; content_type=%s; content_length=%s; "
        "transfer_encoding=%s; received_bytes=%d",
        ", ".join(fields),
        _request.headers.get("content-type", "<none>"),
        _request.headers.get("content-length", "<none>"),
        _request.headers.get("transfer-encoding", "<none>"),
        received_bytes,
    )
    return JSONResponse(
        status_code=422,
        content={"detail": "请求字段校验失败", "fields": fields},
    )


class AnnotationRequest(BaseModel):
    lesson_id: str = Field(min_length=1, max_length=64)
    version_id: str = Field(min_length=1, max_length=64)
    subject: str = Field(min_length=1, max_length=100)
    grade: str = Field(min_length=1, max_length=100)
    topic: str = Field(min_length=1, max_length=255)
    content: str = Field(min_length=1, max_length=100000)
    # v1 平台默认生成改进建议；保留该字段便于本地调试时临时关闭。
    suggest: bool = DEFAULT_SUGGEST


class BatchAnnotationRequest(BaseModel):
    items: list[AnnotationRequest] = Field(min_length=2, max_length=20)


class ComparisonRequest(BaseModel):
    run_ids: list[str] = Field(min_length=2, max_length=20)


def _run_dir(run_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", run_id):
        raise HTTPException(404, "任务不存在")
    path = RUN_ROOT / run_id
    if not path.is_dir():
        raise HTTPException(404, "任务不存在")
    return path


def _write_state(run_dir: Path, payload: dict) -> None:
    target = run_dir / "state.json"
    temp = run_dir / "state.tmp"
    with LOCK:
        temp.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temp.replace(target)


def _read_state(run_dir: Path) -> dict:
    return json.loads((run_dir / "state.json").read_text(encoding="utf-8"))


def _artifact_urls(run_id: str, artifacts: dict) -> dict:
    return {
        kind: f"/internal/v1/runs/{run_id}/artifacts/{kind}"
        for kind in artifacts
    }


def _artifact_entries(run_id: str, artifacts: dict) -> list[dict]:
    """返回新版前端更易消费的产物列表，同时保留旧 artifacts 映射。"""
    return [
        {
            "kind": kind,
            "name": Path(relative_path).name,
            "url": f"/internal/v1/runs/{run_id}/artifacts/{kind}",
        }
        for kind, relative_path in artifacts.items()
    ]


def _public_state(run_id: str, record: dict) -> dict:
    """对平台返回 HTTP 可用的产物入口，不暴露本机文件系统相对路径。"""
    payload = dict(record)
    stored = record.get("artifacts", {})
    if isinstance(stored, dict):
        payload["artifacts"] = _artifact_urls(run_id, stored)
    return payload


def aggregate_student_rubric(result: dict) -> dict:
    """把 v1.2 的 12 个学生标准细项稳定转换为平台 rubric 合同。

    兼容目标：
    - 保留 rubric.groups，避免当前统一 Web 立即报错；
    - 新增 rubric.dimensions，供十二维雷达图与二级维度卡片直接使用；
    - A/B/C 由同一套学生标准分值区间统一换算，不信任模型自由生成等级；
    - 同步修正 rawResult 中的维度等级、批注评价与总分，确保导出/接口口径一致。
    """
    rows = result.get("维度评分")
    if not isinstance(rows, list):
        raise ValueError("F1 未返回维度评分")

    by_name: dict[str, dict] = {}
    for row in rows:
        name = row.get("维度名称")
        if name in by_name or name not in EXPECTED:
            raise ValueError(f"F1 评分细项重复或未知：{name}")
        by_name[name] = row

    if set(by_name) != EXPECTED:
        raise ValueError(f"F1 评分细项缺失：{sorted(EXPECTED - set(by_name))}")

    dimensions = []
    groups = []
    for name, maximum, children in GROUPS:
        items = [by_name[child] for child in children]
        full_marks = [int(item["满分"]) for item in items]
        if sum(full_marks) != maximum:
            raise ValueError(f"{name} 满分与学生标准不一致")

        points = [float(item["得分"]) for item in items]
        if any(score < 0 or score > cap for score, cap in zip(points, full_marks)):
            raise ValueError(f"{name} 得分超出细项范围")

        normalized_items = []
        for item, full_mark, score in zip(items, full_marks, points):
            grade = score_to_grade(full_mark, score)
            if grade not in {"A", "B", "C"}:
                raise ValueError(f"{item.get('维度名称')} 等级换算失败")

            # API 与 v1.2 Word/JSON 产物必须使用同一评价口径。
            item["等级"] = grade
            for annotation in item.get("批注列表", []) or []:
                if isinstance(annotation, dict):
                    annotation["评价"] = grade

            normalized_items.append(item)
            dimensions.append(
                {
                    "name": item["维度名称"],
                    "score": score,
                    "maximum": full_mark,
                    "grade": grade,
                }
            )

        groups.append(
            {
                "name": name,
                "score": sum(points),
                "maximum": maximum,
                # 保留 v1.1 的 items 形态，旧 Web 仍可读取原始中文字段。
                "items": normalized_items,
            }
        )

    total = sum(item["score"] for item in dimensions)
    summary = result.setdefault("评分汇总", {})
    summary["总分"] = total
    summary["各维度得分"] = {item["name"]: item["score"] for item in dimensions}

    return {
        "standard": "F1-student-V3",
        "score": total,
        "maximum": 100,
        "groups": groups,
        "dimensions": dimensions,
        "gradingScale": {
            "A": "达到该维度较高要求",
            "B": "基本达到要求，仍需改进",
            "C": "存在明显不足",
        },
    }


def _initialize_run(request: AnnotationRequest) -> tuple[str, Path, dict]:
    if not request.content.strip():
        raise HTTPException(422, "教案正文不能为空")
    run_id = uuid.uuid4().hex
    run_dir = RUN_ROOT / run_id
    run_dir.mkdir(parents=True, exist_ok=False)
    record = {
        "runId": run_id,
        "lessonId": request.lesson_id,
        "versionId": request.version_id,
        "subject": request.subject,
        "grade": request.grade,
        "topic": request.topic,
        "status": "QUEUED",
        "criteriaType": CRITERIA_TYPE,
        "variant": PROMPT_VARIANT,
        "suggest": request.suggest,
        "createdAt": datetime.now(timezone.utc).isoformat(),
    }
    _write_state(run_dir, record)
    return run_id, run_dir, record


def _execute(run_id: str, request: AnnotationRequest) -> None:
    run_dir = RUN_ROOT / run_id
    record = _read_state(run_dir)
    record["status"] = "RUNNING"
    _write_state(run_dir, record)

    try:
        # F2/主平台传的是已保存版本的正文文本；F1 原流程以文档为输入，
        # 因此在当前任务目录中重建一个临时 DOCX，再交给 run_annotation()。
        input_dir = run_dir / "input"
        input_dir.mkdir(parents=True, exist_ok=True)
        filename = "lesson.docx"

        doc = Document()
        doc.add_heading(request.topic, level=1)
        for line in (f"学科：{request.subject}", f"年级：{request.grade}"):
            doc.add_paragraph(line)
        for line in request.content.splitlines():
            doc.add_paragraph(line)
        doc.save(input_dir / filename)

        config = deepcopy(load_config(ROOT / "config.yaml"))
        config.setdefault("data", {})["plans_dir"] = str(input_dir)
        config.setdefault("output", {})["dir"] = str(run_dir / "output")

        # 平台整合口径：学生标准 + V2 引导提示词。
        outcome = run_annotation(
            str(ROOT),
            config,
            filename,
            criteria_type=CRITERIA_TYPE,
            variant=PROMPT_VARIANT,
            suggest=request.suggest,
        )

        raw_result = outcome["result"]
        rubric = aggregate_student_rubric(raw_result)

        # 只接纳本次 run_dir 之内的真实文件，避免把任意路径暴露给下载接口。
        stored_artifacts: dict[str, str] = {}
        for kind, raw_path in outcome.get("files", {}).items():
            path = Path(raw_path).resolve()
            if path.is_file() and path.is_relative_to(run_dir.resolve()):
                stored_artifacts[kind] = str(path.relative_to(run_dir))

        artifact_urls = _artifact_urls(run_id, stored_artifacts)
        result_payload = {
            "runId": run_id,
            "lessonId": request.lesson_id,
            "versionId": request.version_id,
            "subject": request.subject,
            "grade": request.grade,
            "topic": request.topic,
            "criteriaType": CRITERIA_TYPE,
            "variant": PROMPT_VARIANT,
            "suggest": request.suggest,
            "status": "COMPLETED",
            "score": {
                "total": rubric["score"],
                "maximum": rubric["maximum"],
            },
            "summary": raw_result.get("总体评价", ""),
            "rubric": rubric,
            "rawResult": raw_result,
            # 新合同使用 files 列表；artifacts 映射继续保留，避免旧 Web 失效。
            "files": _artifact_entries(run_id, stored_artifacts),
            "artifacts": artifact_urls,
            "elapsedSeconds": outcome["elapsed"],
        }
        (run_dir / "result.json").write_text(
            json.dumps(result_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        record.update(
            status="COMPLETED",
            artifacts=stored_artifacts,
            finishedAt=datetime.now(timezone.utc).isoformat(),
        )
    except Exception:
        LOGGER.exception("F1 annotation failed for run %s", run_id)
        record.update(
            status="FAILED",
            error="批注任务失败；请查看 F1 接口终端日志",
            finishedAt=datetime.now(timezone.utc).isoformat(),
        )

    _write_state(run_dir, record)


@app.get("/health", include_in_schema=False)
@app.get("/internal/v1/health")
def health():
    config = load_config(ROOT / "config.yaml")
    provider = config.get("llm", {}).get("provider", "deepseek")
    configured = get_key_status(
        config.get("api_keys", {}).get(provider, ""),
        provider,
    )["configured"]
    return {
        "status": "up",
        "rubric": "F1-student-V3",
        "criteriaType": CRITERIA_TYPE,
        "variant": PROMPT_VARIANT,
        "suggestDefault": DEFAULT_SUGGEST,
        "mode": "real",
        "model_configured": configured,
        "capabilities": {
            "singleAnnotation": True,
            "batchAnnotation": True,
            "comparison": True,
            "batchQualityReport": True,
        },
    }


@app.post("/internal/v1/runs/annotate", status_code=202)
def annotate(request: AnnotationRequest):
    if not request.content.strip():
        LOGGER.warning("F1 blank lesson content: chars=%d", len(request.content))
    run_id, _run_dir_path, record = _initialize_run(request)
    EXECUTOR.submit(_execute, run_id, request)
    return _public_state(run_id, record)


@app.get("/internal/v1/runs/{run_id}")
def state(run_id: str):
    record = _read_state(_run_dir(run_id))
    return _public_state(run_id, record)


@app.get("/internal/v1/runs/{run_id}/result")
def result(run_id: str):
    run_dir = _run_dir(run_id)
    record = _read_state(run_dir)
    if record["status"] == "FAILED":
        raise HTTPException(409, record.get("error", "任务执行失败"))
    if record["status"] != "COMPLETED":
        raise HTTPException(409, "结果尚未完成")
    return json.loads((run_dir / "result.json").read_text(encoding="utf-8"))


@app.get("/internal/v1/runs/{run_id}/artifacts/{kind}")
def artifact(run_id: str, kind: str):
    run_dir = _run_dir(run_id)
    record = _read_state(run_dir)
    stored = record.get("artifacts", {}).get(kind)
    if record["status"] != "COMPLETED" or not stored:
        raise HTTPException(404, "成果不存在")

    path = (run_dir / stored).resolve()
    if not path.is_relative_to(run_dir.resolve()) or not path.is_file():
        raise HTTPException(404, "成果不存在")
    return FileResponse(path, filename=path.name)

# ============================================================
# Batch annotation / comparison platform API
# ============================================================

def _job_dir(root: Path, job_id: str, missing_message: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", job_id):
        raise HTTPException(404, missing_message)
    path = root / job_id
    if not path.is_dir():
        raise HTTPException(404, missing_message)
    return path


def _job_artifact_urls(prefix: str, job_id: str, artifacts: dict) -> dict:
    return {
        kind: f"/internal/v1/{prefix}/{job_id}/artifacts/{kind}"
        for kind in artifacts
    }


def _write_json(path: Path, payload: dict) -> None:
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def _batch_public_state(batch_id: str, record: dict) -> dict:
    payload = dict(record)
    artifacts = record.get("artifacts") or {}
    payload["artifacts"] = _job_artifact_urls("batches", batch_id, artifacts)
    return payload


def _comparison_public_state(comparison_id: str, record: dict) -> dict:
    payload = dict(record)
    artifacts = record.get("artifacts") or {}
    payload["artifacts"] = _job_artifact_urls("comparisons", comparison_id, artifacts)
    return payload


def _execute_batch(batch_id: str, request: BatchAnnotationRequest) -> None:
    batch_dir = BATCH_ROOT / batch_id
    record = _read_state(batch_dir)
    record.update(status="RUNNING", startedAt=datetime.now(timezone.utc).isoformat())
    _write_state(batch_dir, record)

    success_results: list[dict] = []
    item_records: list[dict] = []
    failed_count = 0

    try:
        for index, item in enumerate(request.items, start=1):
            try:
                run_id, _run_dir_path, run_record = _initialize_run(item)
                current = {
                    "index": index,
                    "runId": run_id,
                    "lessonId": item.lesson_id,
                    "versionId": item.version_id,
                    "subject": item.subject,
                    "grade": item.grade,
                    "topic": item.topic,
                    "status": "RUNNING",
                }
                item_records.append(current)
                record.update(
                    completed=index - 1,
                    failed=failed_count,
                    items=item_records,
                )
                _write_state(batch_dir, record)

                # 批量任务在同一工作线程内同步复用现有单份执行逻辑，避免产生第二套评分实现。
                _execute(run_id, item)
                final_run_state = _read_state(RUN_ROOT / run_id)
                current["status"] = final_run_state.get("status")
                if final_run_state.get("status") == "COMPLETED":
                    run_result = json.loads((RUN_ROOT / run_id / "result.json").read_text(encoding="utf-8"))
                    current["score"] = (run_result.get("score") or {}).get("total")
                    success_results.append(run_result)
                else:
                    failed_count += 1
                    current["error"] = final_run_state.get("error", "F1 单份评审失败")
            except Exception:
                failed_count += 1
                LOGGER.exception("F1 batch %s item %s failed", batch_id, index)
                if len(item_records) < index:
                    item_records.append(
                        {
                            "index": index,
                            "runId": None,
                            "lessonId": item.lesson_id,
                            "versionId": item.version_id,
                            "subject": item.subject,
                            "grade": item.grade,
                            "topic": item.topic,
                            "status": "FAILED",
                            "error": "批量任务中的单份评审初始化失败",
                        }
                    )
                else:
                    item_records[-1]["status"] = "FAILED"
                    item_records[-1]["error"] = "批量任务中的单份评审失败"

            record.update(
                completed=index,
                succeeded=len(success_results),
                failed=failed_count,
                items=item_records,
            )
            _write_state(batch_dir, record)

        stored_artifacts: dict[str, str] = {}
        artifacts_dir = batch_dir / "artifacts"
        artifacts_dir.mkdir(parents=True, exist_ok=True)

        if success_results:
            quality_files = write_batch_quality_report(success_results, artifacts_dir)
            for kind, raw_path in quality_files.items():
                path = Path(raw_path).resolve()
                stored_artifacts[kind] = str(path.relative_to(batch_dir.resolve()))

        if len(success_results) >= 2:
            comparison_payload = build_comparison_payload(success_results)
            comparison_json = artifacts_dir / "batch_comparison.json"
            comparison_json.write_text(
                json.dumps(comparison_payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            stored_artifacts["comparison_json"] = str(comparison_json.relative_to(batch_dir))

            legacy_results = []
            for run_result in success_results:
                label = (
                    run_result.get("topic")
                    or run_result.get("lessonId")
                    or run_result.get("runId")
                )
                version = run_result.get("versionId")
                if version:
                    label = f"{label} [{version}]"
                legacy_results.append((label, run_result.get("rawResult") or {}))
            md_path, csv_path = generate_comparison(
                legacy_results,
                str(artifacts_dir),
                CRITERIA_TYPE,
                "F1-student-V3",
            )
            if md_path:
                stored_artifacts["comparison_md"] = str(Path(md_path).resolve().relative_to(batch_dir.resolve()))
            if csv_path:
                stored_artifacts["comparison_csv"] = str(Path(csv_path).resolve().relative_to(batch_dir.resolve()))

        if success_results and failed_count == 0:
            final_status = "COMPLETED"
        elif success_results:
            final_status = "COMPLETED_WITH_ERRORS"
        else:
            final_status = "FAILED"

        result_payload = {
            "batchId": batch_id,
            "status": final_status,
            "total": len(request.items),
            "succeeded": len(success_results),
            "failed": failed_count,
            "items": item_records,
            "artifacts": _job_artifact_urls("batches", batch_id, stored_artifacts),
        }
        if success_results:
            report_payload = json.loads((artifacts_dir / "batch_quality_report.json").read_text(encoding="utf-8"))
            result_payload["qualitySummary"] = report_payload.get("summary")
            result_payload["strengths"] = report_payload.get("strengths")
            result_payload["commonIssues"] = report_payload.get("commonIssues")

        _write_json(batch_dir / "result.json", result_payload)
        record.update(
            status=final_status,
            completed=len(request.items),
            succeeded=len(success_results),
            failed=failed_count,
            items=item_records,
            artifacts=stored_artifacts,
            finishedAt=datetime.now(timezone.utc).isoformat(),
        )
    except Exception:
        LOGGER.exception("F1 batch annotation failed for batch %s", batch_id)
        record.update(
            status="FAILED",
            failed=max(failed_count, 1),
            error="批量评审任务失败；请查看 F1 接口终端日志",
            finishedAt=datetime.now(timezone.utc).isoformat(),
        )

    _write_state(batch_dir, record)


@app.post("/internal/v1/batches/annotate", status_code=202)
def batch_annotate(request: BatchAnnotationRequest):
    blank_items = [index + 1 for index, item in enumerate(request.items) if not item.content.strip()]
    if blank_items:
        raise HTTPException(422, f"以下批量条目教案正文为空：{blank_items}")

    batch_id = uuid.uuid4().hex
    batch_dir = BATCH_ROOT / batch_id
    batch_dir.mkdir(parents=True, exist_ok=False)
    record = {
        "batchId": batch_id,
        "status": "QUEUED",
        "total": len(request.items),
        "completed": 0,
        "succeeded": 0,
        "failed": 0,
        "items": [],
        "createdAt": datetime.now(timezone.utc).isoformat(),
    }
    _write_state(batch_dir, record)
    EXECUTOR.submit(_execute_batch, batch_id, request)
    return _batch_public_state(batch_id, record)


@app.get("/internal/v1/batches/{batch_id}")
def batch_state(batch_id: str):
    batch_dir = _job_dir(BATCH_ROOT, batch_id, "批量任务不存在")
    return _batch_public_state(batch_id, _read_state(batch_dir))


@app.get("/internal/v1/batches/{batch_id}/result")
def batch_result(batch_id: str):
    batch_dir = _job_dir(BATCH_ROOT, batch_id, "批量任务不存在")
    record = _read_state(batch_dir)
    if record.get("status") == "FAILED" and not (batch_dir / "result.json").is_file():
        raise HTTPException(409, record.get("error", "批量任务执行失败"))
    if record.get("status") not in {"COMPLETED", "COMPLETED_WITH_ERRORS", "FAILED"}:
        raise HTTPException(409, "批量结果尚未完成")
    path = batch_dir / "result.json"
    if not path.is_file():
        raise HTTPException(409, "批量任务未产生可用结果")
    return json.loads(path.read_text(encoding="utf-8"))


@app.get("/internal/v1/batches/{batch_id}/artifacts/{kind}")
def batch_artifact(batch_id: str, kind: str):
    batch_dir = _job_dir(BATCH_ROOT, batch_id, "批量任务不存在")
    record = _read_state(batch_dir)
    stored = (record.get("artifacts") or {}).get(kind)
    if record.get("status") not in {"COMPLETED", "COMPLETED_WITH_ERRORS"} or not stored:
        raise HTTPException(404, "批量成果不存在")
    path = (batch_dir / stored).resolve()
    if not path.is_relative_to(batch_dir.resolve()) or not path.is_file():
        raise HTTPException(404, "批量成果不存在")
    return FileResponse(path, filename=path.name)


@app.post("/internal/v1/comparisons", status_code=201)
def create_comparison(request: ComparisonRequest):
    run_ids = list(dict.fromkeys(request.run_ids))
    if len(run_ids) < 2:
        raise HTTPException(422, "至少需要 2 个不同的 runId")

    results: list[dict] = []
    for run_id in run_ids:
        run_dir = _run_dir(run_id)
        record = _read_state(run_dir)
        if record.get("status") != "COMPLETED":
            raise HTTPException(409, f"runId {run_id} 尚未完成，不能参与对比")
        results.append(json.loads((run_dir / "result.json").read_text(encoding="utf-8")))

    comparison_id = uuid.uuid4().hex
    comparison_dir = COMPARISON_ROOT / comparison_id
    artifacts_dir = comparison_dir / "artifacts"
    artifacts_dir.mkdir(parents=True, exist_ok=False)

    payload = build_comparison_payload(results)
    payload.update(comparisonId=comparison_id, status="COMPLETED", runIds=run_ids)
    _write_json(comparison_dir / "result.json", payload)

    json_path = artifacts_dir / "comparison.json"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    legacy_results = []
    for run_result in results:
        label = run_result.get("topic") or run_result.get("lessonId") or run_result.get("runId")
        if run_result.get("versionId"):
            label = f"{label} [{run_result['versionId']}]"
        legacy_results.append((label, run_result.get("rawResult") or {}))
    md_path, csv_path = generate_comparison(
        legacy_results,
        str(artifacts_dir),
        CRITERIA_TYPE,
        "F1-student-V3",
    )

    stored_artifacts = {"json": str(json_path.relative_to(comparison_dir))}
    if md_path:
        stored_artifacts["md"] = str(Path(md_path).resolve().relative_to(comparison_dir.resolve()))
    if csv_path:
        stored_artifacts["csv"] = str(Path(csv_path).resolve().relative_to(comparison_dir.resolve()))

    record = {
        "comparisonId": comparison_id,
        "status": "COMPLETED",
        "runIds": run_ids,
        "count": len(run_ids),
        "artifacts": stored_artifacts,
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "finishedAt": datetime.now(timezone.utc).isoformat(),
    }
    _write_state(comparison_dir, record)
    return _comparison_public_state(comparison_id, record)


@app.get("/internal/v1/comparisons/{comparison_id}")
def comparison_state(comparison_id: str):
    comparison_dir = _job_dir(COMPARISON_ROOT, comparison_id, "对比任务不存在")
    return _comparison_public_state(comparison_id, _read_state(comparison_dir))


@app.get("/internal/v1/comparisons/{comparison_id}/result")
def comparison_result(comparison_id: str):
    comparison_dir = _job_dir(COMPARISON_ROOT, comparison_id, "对比任务不存在")
    return json.loads((comparison_dir / "result.json").read_text(encoding="utf-8"))


@app.get("/internal/v1/comparisons/{comparison_id}/artifacts/{kind}")
def comparison_artifact(comparison_id: str, kind: str):
    comparison_dir = _job_dir(COMPARISON_ROOT, comparison_id, "对比任务不存在")
    record = _read_state(comparison_dir)
    stored = (record.get("artifacts") or {}).get(kind)
    if record.get("status") != "COMPLETED" or not stored:
        raise HTTPException(404, "对比成果不存在")
    path = (comparison_dir / stored).resolve()
    if not path.is_relative_to(comparison_dir.resolve()) or not path.is_file():
        raise HTTPException(404, "对比成果不存在")
    return FileResponse(path, filename=path.name)
