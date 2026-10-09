"""
输出生成模块
生成带批注的 DOCX、JSON 评分清单、Markdown 报告
"""

import datetime
import json
import os
from typing import Dict

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# 颜色
C_GREEN = RGBColor(0x27, 0xAE, 0x60)    # 达标
C_PURPLE = RGBColor(0x8E, 0x44, 0xAD)   # 优秀/亮点
C_ORANGE = RGBColor(0xF3, 0x9C, 0x12)   # 部分达标
C_RED = RGBColor(0xE7, 0x4C, 0x3C)      # 未达标
C_BLUE = RGBColor(0x29, 0x80, 0xB9)     # 基本达标
C_GRAY = RGBColor(0x7F, 0x8C, 0x8D)     # 默认

STATUS_COLORS = {
    # 2026-10-04：等级改为 A/B/C（按学生标准分值区间换算）
    "A": C_GREEN,        # 达标（A 档）
    "B": C_ORANGE,       # 需改进（B 档）
    "C": C_RED,          # 明显不足（C 档）
    # 兼容历史产物中的五档写法
    "优秀": C_PURPLE, "亮点": C_PURPLE,
    "达标": C_GREEN,
    "基本达标": C_BLUE,
    "部分达标": C_ORANGE,
    "未达标": C_RED, "不达标": C_RED,
}


def _add_annotation_block(doc, annotation: Dict):
    """添加一个带左边框的批注卡片"""
    status = annotation.get("评价", "")
    color = STATUS_COLORS.get(status, C_GRAY)

    # 段落 + 左边框
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '12')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), str(color))
    pBdr.append(left)
    pPr.append(pBdr)

    badge = p.add_run(f"批注 | {status}")
    badge.bold = True
    badge.font.color.rgb = color
    badge.font.size = Pt(9)

    # 位置
    p2 = doc.add_paragraph()
    r = p2.add_run(f"位置：{annotation.get('位置', '未标注')}")
    r.font.size = Pt(9)
    r.font.color.rgb = C_GRAY

    # 原文
    p3 = doc.add_paragraph()
    ql = p3.add_run("原文：")
    ql.font.size = Pt(9)
    ql.font.color.rgb = C_GRAY
    qt = p3.add_run(f"「{annotation.get('原文引用', '无')}」")
    qt.font.size = Pt(9)
    qt.italic = True

    # 分析
    p4 = doc.add_paragraph()
    a = p4.add_run(f"分析：{annotation.get('具体分析', '无')}")
    a.font.size = Pt(10)

    # 建议（纯诊断）
    sug = annotation.get('建议', '') or annotation.get('改进建议', '')
    if sug and sug != '无需改进':
        p5 = doc.add_paragraph()
        s = p5.add_run(f"建议（诊断）：{sug}")
        s.font.size = Pt(10)
        s.font.color.rgb = C_BLUE

    # 依据来源（新增）
    basis = annotation.get('依据来源', '')
    if basis:
        p6 = doc.add_paragraph()
        t = p6.add_run(f"\U0001F4DA 依据来源：{basis}")
        t.font.size = Pt(9)
        t.font.color.rgb = RGBColor(0x8E, 0x44, 0xAD)

    # 教学法参考
    method_ref = annotation.get('教学法参考', '')
    if method_ref:
        p6b = doc.add_paragraph()
        tm = p6b.add_run(f"\U0001F3EB 教学法参考：{method_ref}")
        tm.font.size = Pt(9)
        tm.font.color.rgb = RGBColor(0x00, 0x8A, 0x00)

    # 细化建议（具体修改方法）
    detail = annotation.get('细化建议', '')
    if detail:
        p7 = doc.add_paragraph()
        d = p7.add_run(f"\U0001F4A1 细化建议：{detail}")
        d.font.size = Pt(10)
        d.font.color.rgb = C_PURPLE


def _setup_font(doc):
    """设置默认字体为微软雅黑"""
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '微软雅黑')
    rPr.append(rFonts)


def generate_docx(result: Dict, output_path: str) -> str:
    """生成带批注卡片的分页 DOCX 报告"""
    doc = Document()
    _setup_font(doc)

    # 封面
    doc.add_heading("教学设计批注与评分报告", level=0)
    info = result.get("教案基本信息", {})
    p = doc.add_paragraph()
    r = p.add_run(f"课题：{info.get('课题名称','')} | 学科：{info.get('学科','')} | "
                  f"年级：{info.get('年级','')} | 课型：{info.get('课型','')}")
    r.font.size = Pt(9); r.font.color.rgb = C_GRAY

    p = doc.add_paragraph()
    r = p.add_run(f"评审日期：{datetime.datetime.now().strftime('%Y-%m-%d')} | AI 智能评审")
    r.font.size = Pt(9); r.font.color.rgb = C_GRAY

    doc.add_paragraph()

    # 评分总表（2026-10-04：按学生标准的一级维度分组 + 标注 A/B/C 等级）
    summary = result.get("评分汇总", {})
    dim_max = sum(d.get("满分", 0) for d in result.get("维度评分", []))
    total = summary.get('总分', '?')
    doc.add_heading(f"总分：{total} / {dim_max}", level=1)

    from standards import group_by_level1
    groups = group_by_level1(result.get("维度评分", []))

    table = doc.add_table(rows=1, cols=5)
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(['一级维度', '二级维度', '满分', '得分', '等级']):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True

    for g in groups:
        # 一级维度小计行
        row = table.add_row()
        row.cells[0].text = g["一级维度"]
        row.cells[1].text = "（小计）"
        row.cells[2].text = str(g["满分"])
        row.cells[3].text = str(g["得分"])
        row.cells[4].text = g.get("等级", "")
        for c in row.cells:
            for pp in c.paragraphs:
                for rr in pp.runs:
                    rr.bold = True
        # 二级维度明细
        for d in g["维度列表"]:
            row = table.add_row()
            row.cells[0].text = ""
            row.cells[1].text = d.get("维度名称", "")
            row.cells[2].text = str(d.get("满分", 0))
            row.cells[3].text = str(d.get("得分", 0))
            row.cells[4].text = d.get("等级", "")

    doc.add_paragraph()

    # 总体评价
    doc.add_heading("总体评价", level=1)
    p = doc.add_paragraph(result.get("总体评价", ""))

    # 逐维度（按一级维度分组：一级维度作章节，二级维度作小节）
    doc.add_heading("逐维度详细批注", level=1)
    for g in groups:
        doc.add_heading(f"{g['一级维度']}（{g['得分']}/{g['满分']} 分）", level=1)
        for dim in g["维度列表"]:
            grade = dim.get("等级", "")
            doc.add_heading(
                f"{dim.get('维度名称','')}：{dim.get('得分',0)}/{dim.get('满分',0)} 分"
                f"{'（等级 ' + grade + '）' if grade else ''}", level=2)
            p = doc.add_paragraph(dim.get("得分理由", ""))
            if p.runs:
                p.runs[0].font.color.rgb = C_GRAY

            for ann in dim.get("批注列表", []):
                _add_annotation_block(doc, ann)

        doc.add_page_break()

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    try:
        doc.save(output_path)
    except PermissionError:
        import time
        alt = output_path.replace(".docx", f"_{int(time.time())}.docx")
        doc.save(alt)
        print(f"  ⚠ 原文件被锁定，已保存到: {alt}")
    return output_path


def generate_json(result: Dict, output_path: str) -> str:
    """生成 JSON 评分清单"""
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    return output_path


def generate_markdown(result: Dict, output_path: str) -> str:
    """生成 Markdown 评分报告"""
    info = result.get("教案基本信息", {})
    summary = result.get("评分汇总", {})

    md = "# 教学设计评分清单\n\n"
    md += f"**课题**：{info.get('课题名称','')} | **学科**：{info.get('学科','')} | "
    md += f"**年级**：{info.get('年级','')} | **课型**：{info.get('课型','')}\n"
    md += f"**评审日期**：{datetime.datetime.now().strftime('%Y-%m-%d')} | "
    dim_max = sum(d.get("满分", 0) for d in result.get("维度评分", []))
    md += f"**总分**：{summary.get('总分','?')}/{dim_max}\n\n---\n\n"

    md += "## 总体评价\n\n" + result.get("总体评价","") + "\n\n---\n\n"

    md += "## 逐维度批注\n\n"
    for dim in result.get("维度评分", []):
        _g = dim.get("等级") or ""
        md += f"### {dim.get('维度名称','')}：{dim.get('得分',0)}/{dim.get('满分',0)}分{f'（等级 {_g}）' if _g else ''}\n"
        md += f"> {dim.get('得分理由','')}\n\n"
        for i, ann in enumerate(dim.get("批注列表", []), 1):
            emoji = {"A": "✅", "B": "⚠️", "C": "❌",
                     "达标": "✅", "部分达标": "⚠️", "未达标": "❌"}.get(ann.get('评价', ''), "")
            md += f"**批注 {i}** {emoji} {ann.get('评价','')}\n"
            md += f"- 位置：{ann.get('位置','')}\n"
            md += f"- 原文：> {ann.get('原文引用','')}\n"
            md += f"- 分析：{ann.get('具体分析','')}\n"
            sug = ann.get('建议','') or ann.get('改进建议','')
            if sug and sug != '无需改进':
                md += f"- 🔍 建议（诊断）：{sug}\n"
            basis = ann.get('依据来源','')
            if basis:
                md += f"- 📚 依据来源：{basis}\n"
            mref = ann.get('教学法参考','')
            if mref:
                md += f"- 🏫 教学法参考：{mref}\n"
            detail = ann.get('细化建议','')
            if detail:
                md += f"- 💡 细化建议：{detail}\n"
            md += "\n"
        md += "---\n\n"

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(md)
    return output_path
